"""Batches ("items" in the frontend). A batch is created under the product that owns its barcode;
the product is created on first sight and its name, prices and supplier follow the latest entry."""
import csv
import io
import re
import uuid
from datetime import date, datetime

from fastapi import APIRouter, Depends, File, HTTPException, Query, Response, UploadFile, status
from pydantic import ValidationError
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from ..db import get_db
from ..deps import current_user, require
from ..models import Batch, Product, User
from ..schemas import ImportResult, ImportRowError, ItemCreate, ItemOut, ItemUpdate, ParseResult
from ..serializers import item_out, sold_this_month

router = APIRouter(prefix="/items", tags=["items"])


async def _get_batch(db: AsyncSession, user: User, item_id: str) -> Batch:
    try:
        bid = uuid.UUID(item_id)
    except ValueError:
        raise HTTPException(404, "That batch doesn't exist.")
    b = await db.scalar(select(Batch).join(Product).where(Batch.id == bid, Product.pharmacy_id == user.pharmacy_id))
    if not b:
        raise HTTPException(404, "That batch doesn't exist.")
    return b


@router.get("", response_model=list[ItemOut])
async def list_items(user: User = Depends(require("view_inventory")), db: AsyncSession = Depends(get_db)):
    batches = (await db.scalars(
        select(Batch).join(Product).where(Product.pharmacy_id == user.pharmacy_id).order_by(Batch.expiry_date)
    )).unique().all()
    sold = await sold_this_month(db, [b.id for b in batches])
    return [item_out(b, sold.get(b.id, 0)) for b in batches]


@router.get("/barcode/{code}", response_model=list[ItemOut])
async def by_barcode(code: str, user: User = Depends(current_user), db: AsyncSession = Depends(get_db)):
    """Every batch of the product with this barcode, first-expiry-first-out."""
    batches = (await db.scalars(
        select(Batch).join(Product).where(Product.pharmacy_id == user.pharmacy_id, Product.barcode == code.strip()).order_by(Batch.expiry_date)
    )).unique().all()
    return [item_out(b) for b in batches]


@router.post("", response_model=ItemOut, status_code=201)
async def create_item(body: ItemCreate, user: User = Depends(require("edit_inventory")), db: AsyncSession = Depends(get_db)):
    product = await db.scalar(select(Product).where(Product.pharmacy_id == user.pharmacy_id, Product.barcode == body.barcode))
    if product is None:
        product = Product(pharmacy_id=user.pharmacy_id, barcode=body.barcode, name=body.name, category=body.category,
                          supplier=body.supplier, unit_cost=body.unit_cost, unit_price=body.unit_price,
                          reorder_level=body.reorder_level if body.reorder_level is not None else 20)
        db.add(product)
        await db.flush()
    else:
        product.name, product.category = body.name, body.category or product.category
        product.unit_cost, product.unit_price = body.unit_cost, body.unit_price
        product.supplier = body.supplier or product.supplier
        exists = await db.scalar(select(Batch.id).where(Batch.product_id == product.id, Batch.batch_no.ilike(body.batch)))
        if exists:
            raise HTTPException(status.HTTP_409_CONFLICT, f"Batch {body.batch} already exists for this barcode. Edit that row instead.")
    batch = Batch(product_id=product.id, batch_no=body.batch, expiry_date=body.expiry, qty=body.qty)
    db.add(batch)
    await db.commit()
    await db.refresh(batch, ["product"])
    return item_out(batch)


@router.patch("/{item_id}", response_model=ItemOut)
async def update_item(item_id: str, body: ItemUpdate, user: User = Depends(require("edit_inventory")), db: AsyncSession = Depends(get_db)):
    b = await _get_batch(db, user, item_id)
    data = body.model_dump(exclude_unset=True)
    p = b.product
    for k_api, k_db, target in [("batch", "batch_no", b), ("expiry", "expiry_date", b), ("qty", "qty", b),
                                ("name", "name", p), ("category", "category", p), ("supplier", "supplier", p),
                                ("unit_cost", "unit_cost", p), ("unit_price", "unit_price", p), ("reorder_level", "reorder_level", p)]:
        if k_api in data and data[k_api] is not None:
            setattr(target, k_db, data[k_api])
    if float(p.unit_price) < float(p.unit_cost):
        raise HTTPException(422, "Selling price is below unit cost")
    await db.commit()
    await db.refresh(b, ["product"])
    sold = await sold_this_month(db, [b.id])
    return item_out(b, sold.get(b.id, 0))


@router.delete("/{item_id}", status_code=204)
async def delete_item(item_id: str, user: User = Depends(require("delete_inventory")), db: AsyncSession = Depends(get_db)):
    b = await _get_batch(db, user, item_id)
    await db.delete(b)
    await db.commit()
    return Response(status_code=204)


# ---------------- Import (CSV or Excel) ----------------
ALIASES = {
    "product": "name", "product_name": "name", "item": "name", "drug": "name",
    "batch_no": "batch", "batch_number": "batch", "lot": "batch",
    "expiry_date": "expiry", "exp": "expiry", "expires": "expiry", "expiration": "expiry",
    "quantity": "qty", "units": "qty", "stock": "qty",
    "cost": "unit_cost", "unitcost": "unit_cost", "cost_price": "unit_cost",
    "price": "unit_price", "unitprice": "unit_price", "selling_price": "unit_price",
    "reorder": "reorder_level", "reorderlevel": "reorder_level",
    "sold": "sold_this_month", "units_sold": "sold_this_month", "qty_sold": "sold_this_month", "sales": "sold_this_month",
    "soldthismonth": "sold_this_month",
}
MAX_ROWS = 5000


def _key(h) -> str:
    k = re.sub(r"[\s\-]+", "_", str(h or "").strip().lower())
    return ALIASES.get(k, k)


def _read_rows(filename: str, raw: bytes) -> list[dict]:
    if filename.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        ws = load_workbook(io.BytesIO(raw), read_only=True, data_only=True).active
        it = ws.iter_rows(values_only=True)
        head = [_key(h) for h in next(it, [])]
        return [dict(zip(head, r)) for r in it if any(c not in (None, "") for c in r)]
    if filename.lower().endswith((".csv", ".txt")):
        text = raw.decode("utf-8-sig", errors="replace")
        reader = csv.DictReader(io.StringIO(text))
        return [{_key(k): v for k, v in r.items()} for r in reader if any((v or "").strip() for v in r.values())]
    raise HTTPException(415, "Upload a .csv or .xlsx file.")


def _date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v or "").strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%m/%Y", "%d %b %Y", "%b %Y"):
        try:
            d = datetime.strptime(s, fmt).date()
            return d
        except ValueError:
            continue
    return s  # let validation report it


def _clean(row: dict) -> dict:
    code = row.get("barcode")
    if isinstance(code, float):
        code = f"{code:.0f}"
    return {
        "name": str(row.get("name") or "").strip(), "category": str(row.get("category") or "").strip(),
        "barcode": str(code or "").strip(), "batch": str(row.get("batch") or "").strip(),
        "expiry": _date(row.get("expiry")), "qty": row.get("qty") or 0,
        "unit_cost": row.get("unit_cost") or 0, "unit_price": row.get("unit_price") or 0,
        "supplier": str(row.get("supplier") or "").strip(),
        **({"reorder_level": row["reorder_level"]} if row.get("reorder_level") not in (None, "") else {}),
        **({"sold_this_month": row["sold_this_month"]} if row.get("sold_this_month") not in (None, "") else {}),
    }


async def _upsert(db: AsyncSession, user: User, it: ItemCreate) -> str:
    product = await db.scalar(select(Product).where(Product.pharmacy_id == user.pharmacy_id, Product.barcode == it.barcode))
    if product is None:
        product = Product(pharmacy_id=user.pharmacy_id, barcode=it.barcode, name=it.name, category=it.category, supplier=it.supplier,
                          unit_cost=it.unit_cost, unit_price=it.unit_price, reorder_level=it.reorder_level if it.reorder_level is not None else 20)
        db.add(product)
        await db.flush()
    else:
        product.name, product.unit_cost, product.unit_price = it.name, it.unit_cost, it.unit_price
        product.category = it.category or product.category
        product.supplier = it.supplier or product.supplier
        if it.reorder_level is not None:
            product.reorder_level = it.reorder_level
    batch = await db.scalar(select(Batch).where(Batch.product_id == product.id, Batch.batch_no.ilike(it.batch)))
    if batch:
        batch.expiry_date, batch.qty = it.expiry, it.qty
        return "updated"
    db.add(Batch(product_id=product.id, batch_no=it.batch, expiry_date=it.expiry, qty=it.qty))
    await db.flush()
    return "added"


async def _validate_upload(file: UploadFile) -> tuple[int, list[ItemCreate], list[ImportRowError]]:
    raw = await file.read()
    if len(raw) > 5 * 1024 * 1024:
        raise HTTPException(413, "That file is over 5 MB. Split it into smaller files.")
    rows = _read_rows(file.filename or "", raw)
    if not rows:
        raise HTTPException(400, "The file has no data rows under the header.")
    if len(rows) > MAX_ROWS:
        raise HTTPException(400, f"That file has {len(rows)} rows; the limit is {MAX_ROWS} per file.")
    good: list[ItemCreate] = []
    errors: list[ImportRowError] = []
    for i, r in enumerate(rows, start=2):  # row 1 is the header
        try:
            good.append(ItemCreate.model_validate(_clean(r)))
        except ValidationError as e:
            first = e.errors()[0]
            field = ".".join(str(x) for x in first.get("loc", [])) or "row"
            errors.append(ImportRowError(row=i, error=f"{field}: {str(first.get('msg', 'invalid')).removeprefix('Value error, ')}"))
    return len(rows), good, errors


@router.post("/parse", response_model=ParseResult)
async def parse_file(file: UploadFile = File(...), user: User = Depends(require("view_insights", "import"))):
    """Read and validate a CSV/Excel file without saving it, e.g. for a one-time analysis in AI Insights."""
    total, good, errors = await _validate_upload(file)
    return ParseResult(total=total, rows=good, errors=errors[:200])


@router.post("/import", response_model=ImportResult)
async def import_items(
    file: UploadFile = File(...),
    dry_run: bool = Query(True, description="true = check and preview only; false = save"),
    user: User = Depends(require("import")),
    db: AsyncSession = Depends(get_db),
):
    """Import batches from a CSV or Excel file. Rows that match an existing barcode + batch are updated."""
    total, good, errors = await _validate_upload(file)
    added = updated = 0
    if not dry_run:
        for it in good:
            if await _upsert(db, user, it) == "added":
                added += 1
            else:
                updated += 1
        await db.commit()
    return ImportResult(dry_run=dry_run, total=total, added=added, updated=updated, errors=errors[:200], preview=good[:20])
